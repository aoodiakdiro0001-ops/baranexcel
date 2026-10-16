import os
import sys
import argparse
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


SUPPORTED = (".csv", ".xlsx", ".xls")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def read_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        for enc in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
            try:
                return pd.read_csv(path, encoding=enc, dtype=str)
            except UnicodeDecodeError:
                continue
        raise ValueError("تعذر قراءة ترميز ملف CSV")
    if ext in (".xlsx", ".xls"):
        if ext == ".xls":
            try:
                import xlrd  # noqa: F401
            except ImportError:
                raise ValueError("ملفات .xls تحتاج مكتبة xlrd (pip install xlrd) أو حوّل الملف إلى .xlsx")
        return pd.read_excel(path, dtype=str, engine="openpyxl" if ext == ".xlsx" else None)
    raise ValueError(f"صيغة غير مدعومة: {ext}")


def clean(df):
    report = {}
    report["الصفوف قبل التنظيف"] = len(df)
    report["الأعمدة قبل التنظيف"] = len(df.columns)

    df = df.dropna(axis=1, how="all")
    df = df.dropna(axis=0, how="all")

    df.columns = [str(c).strip() if str(c).strip() else f"عمود_{i+1}" for i, c in enumerate(df.columns)]

    seen = {}
    new_cols = []
    for c in df.columns:
        if c in seen:
            seen[c] += 1
            new_cols.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            new_cols.append(c)
    df.columns = new_cols

    for c in df.columns:
        if df[c].dtype == object:
            df[c] = df[c].astype(str).str.strip()
            df[c] = df[c].replace({"nan": "", "None": "", "NaN": ""})

    # تطبيع الأرقام (1,200 -> 1200) قبل كشف التكرار والا تفلت نسخ مثل 1,200/1200
    for c in df.columns:
        df[c] = to_number(df[c])

    before = len(df)
    df = df.drop_duplicates()
    report["الصفوف المكررة المحذوفة"] = before - len(df)

    df = df.loc[(df != "").any(axis=1)]

    report["الصفوف بعد التنظيف"] = len(df)
    report["الأعمدة بعد التنظيف"] = len(df.columns)
    return df, report


def to_number(series):
    cleaned = series.astype(str).str.replace(",", "", regex=False).str.strip()
    numeric = pd.to_numeric(cleaned, errors="coerce")
    if numeric.notna().mean() > 0.9:
        return numeric
    return series


def write_xlsx(df, out_path, report, source_name):
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="البيانات")
        summary = pd.DataFrame(
            {
                "البيان": ["اسم الملف الأصلي", "تاريخ التنظيف"] + list(report.keys()),
                "القيمة": [source_name, datetime.now().strftime("%Y-%m-%d %H:%M")] + list(report.values()),
            }
        )
        summary.to_excel(writer, index=False, sheet_name="ملخص")

    wb = load_workbook(out_path)
    ws = wb["البيانات"]
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 45)
    wb.save(out_path)


def process(path, out_dir):
    if not os.path.isfile(path):
        print(f"[!] الملف غير موجود: {path}")
        return None
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED:
        print(f"[!] صيغة غير مدعومة: {ext}")
        return None

    print(f"[*] قراءة: {os.path.basename(path)}")
    try:
        df = read_any(path)
    except Exception as ex:
        print(f"[!] فشلت القراءة: {ex}")
        return None
    df, report = clean(df)

    for c in df.columns:
        df[c] = to_number(df[c])

    base = os.path.splitext(os.path.basename(path))[0]
    out_path = os.path.join(out_dir, f"{base}_clean.xlsx")
    write_xlsx(df, out_path, report, os.path.basename(path))

    print(f"[+] تم الإخراج: {out_path}")
    for k, v in report.items():
        print(f"    - {k}: {v}")
    return out_path


def main():
    parser = argparse.ArgumentParser(description="تنظيف ملفات Excel/CSV")
    parser.add_argument("files", nargs="*", help="مسار الملف أو الملفات")
    parser.add_argument("-o", "--out", default=None, help="مجلد الإخراج")
    args = parser.parse_args()

    files = list(args.files)
    if not files:
        raw = input("اسحب الملف هنا أو اكتب مساره ثم Enter: ").strip().strip("'\"")
        if raw:
            files = [raw]

    if not files:
        print("لم يتم إدخال أي ملف.")
        return

    out_dir = args.out or os.path.dirname(os.path.abspath(files[0])) or "."
    os.makedirs(out_dir, exist_ok=True)

    done = 0
    for f in files:
        if process(f, out_dir):
            done += 1
    print(f"\nتمت معالجة {done} ملف/ملفات بنجاح.")


if __name__ == "__main__":
    main()
